from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Dict, List, Tuple

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
DOCS_DIR = ROOT / 'docs'
DIAGRAM_DIR = DOCS_DIR / 'diagrams'
OUTPUT_DOCX = DOCS_DIR / 'eSnagging_Features_Guide.docx'

DOCS_DIR.mkdir(parents=True, exist_ok=True)
DIAGRAM_DIR.mkdir(parents=True, exist_ok=True)


def get_font(size: int, bold: bool = False) -> ImageFont.ImageFont:
    candidates = [
        'arialbd.ttf' if bold else 'arial.ttf',
        'segoeuib.ttf' if bold else 'segoeui.ttf',
    ]
    for name in candidates:
        try:
            return ImageFont.truetype(name, size)
        except OSError:
            continue
    return ImageFont.load_default()


def wrap_text(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.ImageFont, max_width: int) -> List[str]:
    words = text.split()
    if not words:
        return ['']
    lines: List[str] = []
    current = words[0]
    for word in words[1:]:
        candidate = f'{current} {word}'
        width = draw.textbbox((0, 0), candidate, font=font)[2]
        if width <= max_width:
            current = candidate
        else:
            lines.append(current)
            current = word
    lines.append(current)
    return lines


def draw_box(draw: ImageDraw.ImageDraw, xy: Tuple[int, int, int, int], text: str, fill: str = '#EAF2FF') -> None:
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=18, fill=fill, outline='#204170', width=3)
    title_font = get_font(28, bold=True)
    body_font = get_font(22)
    max_width = (x2 - x1) - 24
    lines = wrap_text(draw, text, body_font, max_width)
    line_height = draw.textbbox((0, 0), 'Ag', font=body_font)[3] + 4
    total_h = line_height * len(lines)
    start_y = y1 + ((y2 - y1) - total_h) // 2
    for idx, line in enumerate(lines):
        bbox = draw.textbbox((0, 0), line, font=body_font)
        text_w = bbox[2] - bbox[0]
        draw.text((x1 + ((x2 - x1) - text_w) // 2, start_y + idx * line_height), line, fill='#102543', font=body_font)


def draw_arrow(draw: ImageDraw.ImageDraw, start: Tuple[int, int], end: Tuple[int, int], color: str = '#204170') -> None:
    sx, sy = start
    ex, ey = end
    draw.line((sx, sy, ex, ey), fill=color, width=5)

    if abs(ex - sx) >= abs(ey - sy):
        if ex >= sx:
            points = [(ex, ey), (ex - 16, ey - 10), (ex - 16, ey + 10)]
        else:
            points = [(ex, ey), (ex + 16, ey - 10), (ex + 16, ey + 10)]
    else:
        if ey >= sy:
            points = [(ex, ey), (ex - 10, ey - 16), (ex + 10, ey - 16)]
        else:
            points = [(ex, ey), (ex - 10, ey + 16), (ex + 10, ey + 16)]
    draw.polygon(points, fill=color)


def create_flow_diagram(filename: str, title: str, nodes: Dict[str, Tuple[int, int, int, int, str, str]], edges: List[Tuple[str, str]]) -> Path:
    image = Image.new('RGB', (1800, 1100), color='#F8FBFF')
    draw = ImageDraw.Draw(image)

    title_font = get_font(40, bold=True)
    draw.text((60, 28), title, fill='#0F2D58', font=title_font)

    for _, node in nodes.items():
        x1, y1, x2, y2, label, fill = node
        draw_box(draw, (x1, y1, x2, y2), label, fill=fill)

    for source, target in edges:
        sx1, sy1, sx2, sy2, _, _ = nodes[source]
        tx1, ty1, tx2, ty2, _, _ = nodes[target]

        if sx2 <= tx1:
            start = (sx2, (sy1 + sy2) // 2)
            end = (tx1, (ty1 + ty2) // 2)
        elif tx2 <= sx1:
            start = (sx1, (sy1 + sy2) // 2)
            end = (tx2, (ty1 + ty2) // 2)
        elif sy2 <= ty1:
            start = ((sx1 + sx2) // 2, sy2)
            end = ((tx1 + tx2) // 2, ty1)
        else:
            start = ((sx1 + sx2) // 2, sy1)
            end = ((tx1 + tx2) // 2, ty2)

        draw_arrow(draw, start, end)

    output = DIAGRAM_DIR / filename
    image.save(output)
    return output


def add_title(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(28)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_subtitle(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(11)
    run.italic = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_section(doc: Document, title: str, summary: str) -> None:
    doc.add_heading(title, level=1)
    doc.add_paragraph(summary)


def add_step_table(doc: Document, rows: List[Tuple[str, str, str]]) -> None:
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    hdr = table.rows[0].cells
    hdr[0].text = 'Step / Action'
    hdr[1].text = 'Where to Find It'
    hdr[2].text = 'How to Use It'

    for action, location, usage in rows:
        row = table.add_row().cells
        row[0].text = action
        row[1].text = location
        row[2].text = usage


def add_bullets(doc: Document, items: List[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style='List Bullet')


def build_doc() -> None:
    nav_diagram = create_flow_diagram(
        '01_web_navigation.png',
        'Web Navigation and User Journey',
        {
            'login': (80, 170, 380, 290, 'Login + MFA', '#DDF2FF'),
            'projects': (470, 170, 860, 290, 'Projects Workspace', '#E8F7E8'),
            'project': (950, 170, 1350, 290, 'Project Dashboard', '#EAF2FF'),
            'drawing': (1440, 170, 1740, 290, 'Drawing Viewer', '#FFEEDB'),
            'kanban': (470, 430, 860, 550, 'Kanban Board', '#F2EBFF'),
            'dashboard': (950, 430, 1350, 550, 'Analytics Dashboard', '#F2EBFF'),
            'exports': (1440, 430, 1740, 550, 'Export Center', '#F2EBFF'),
            'quality': (470, 690, 860, 810, 'Templates + Inspections', '#FFF5DE'),
            'ops': (950, 690, 1350, 810, 'Automation + Ops', '#FFF5DE'),
        },
        [
            ('login', 'projects'),
            ('projects', 'project'),
            ('project', 'drawing'),
            ('projects', 'kanban'),
            ('project', 'dashboard'),
            ('dashboard', 'exports'),
            ('projects', 'quality'),
            ('projects', 'ops'),
        ],
    )

    snag_diagram = create_flow_diagram(
        '02_snag_lifecycle.png',
        'Snag Lifecycle With Closeout Guard',
        {
            'new': (80, 220, 320, 340, 'New', '#DDF2FF'),
            'assigned': (400, 220, 700, 340, 'Assigned', '#EAF2FF'),
            'progress': (780, 220, 1080, 340, 'In Progress', '#EAF2FF'),
            'review': (1160, 220, 1510, 340, 'Ready for Review', '#FFF5DE'),
            'closeout': (450, 520, 920, 680, 'Closeout Completion\nMust be 100% unless override', '#FFEAD8'),
            'closed': (1160, 520, 1510, 680, 'Closed', '#E8F7E8'),
            'rejected': (780, 820, 1080, 980, 'Rejected', '#FFE3E3'),
        },
        [
            ('new', 'assigned'),
            ('assigned', 'progress'),
            ('progress', 'review'),
            ('review', 'closeout'),
            ('closeout', 'closed'),
            ('new', 'rejected'),
            ('assigned', 'rejected'),
            ('progress', 'rejected'),
            ('review', 'rejected'),
        ],
    )

    inspection_diagram = create_flow_diagram(
        '03_inspection_approval.png',
        'Inspection Submission and Approval Flow',
        {
            'template': (80, 220, 430, 340, 'Template Builder\n(JSON schema + approval steps)', '#DDF2FF'),
            'draft': (500, 220, 780, 340, 'Draft Submission', '#EAF2FF'),
            'submitted': (860, 220, 1180, 340, 'Submitted / In Review', '#FFF5DE'),
            'approval': (1260, 220, 1720, 340, 'Role-based Approvals\nDecision notes', '#FFF5DE'),
            'signature': (500, 520, 980, 680, 'Digital Signature Upload', '#E8F7E8'),
            'request': (1060, 520, 1520, 680, 'MIR/WIR/IR Request Scheduling', '#E8F7E8'),
            'report': (700, 820, 1160, 980, 'Inspection Reports + Exports', '#F2EBFF'),
        },
        [
            ('template', 'draft'),
            ('draft', 'submitted'),
            ('submitted', 'approval'),
            ('approval', 'signature'),
            ('approval', 'request'),
            ('signature', 'report'),
            ('request', 'report'),
        ],
    )

    mobile_diagram = create_flow_diagram(
        '04_mobile_sync.png',
        'Mobile Offline Sync and Conflict Resolution',
        {
            'offline': (80, 220, 430, 380, 'Offline actions\nCreate/Update snag\nComment/Attach', '#DDF2FF'),
            'queue': (520, 220, 900, 380, 'SQLite Queue + Local Store', '#EAF2FF'),
            'policy': (980, 220, 1380, 380, 'Sync Policy\nInterval + window + throttle', '#FFF5DE'),
            'sync': (1460, 220, 1740, 380, 'Sync Engine\nPush/Pull', '#FFEAD8'),
            'chunk': (520, 560, 900, 720, 'Chunked Attachment Upload', '#E8F7E8'),
            'conflict': (980, 560, 1380, 720, 'Conflict Detected\nLWW + status guard', '#FFE3E3'),
            'resolve': (1460, 560, 1740, 720, 'Resolve Screen\nUse server / Retry / Merge', '#E8F7E8'),
            'done': (980, 860, 1380, 1020, 'Data aligned\nQueue drains', '#F2EBFF'),
        },
        [
            ('offline', 'queue'),
            ('queue', 'policy'),
            ('policy', 'sync'),
            ('queue', 'chunk'),
            ('sync', 'conflict'),
            ('conflict', 'resolve'),
            ('resolve', 'done'),
            ('sync', 'done'),
        ],
    )

    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
    style.font.size = Pt(11)

    add_title(doc, 'eSnagging Feature Guide')
    add_subtitle(doc, f'Generated from the current codebase on {datetime.now().strftime("%Y-%m-%d %H:%M")}.')
    doc.add_paragraph('')

    doc.add_heading('Purpose', level=1)
    doc.add_paragraph(
        'This guide lists the features currently implemented in this repository, where to find each one in the UI, and how to use it step by step. '
        'It is intended for admins, project managers, engineers, consultants, contractors, and field users.'
    )

    doc.add_heading('Quick Access', level=1)
    add_bullets(
        doc,
        [
            'Web login: /login',
            'Main workspace routes: /projects, /board, /dashboard, /exports, /templates, /inspections/*, /equipment, /access-control, /automation, /ops',
            'Mobile tabs: Snags, Equipment, Conflicts, Settings',
            'Demo credentials: admin@sky.demo / password (or any seeded demo user).',
        ],
    )

    doc.add_heading('Diagram 1 - Web Navigation', level=1)
    doc.add_paragraph('Use this as a map of the primary web journey and related modules.')
    doc.add_picture(str(nav_diagram), width=Inches(6.8))

    add_section(
        doc,
        '1. Authentication, Organization Context, and Role-based Visibility',
        'Users authenticate with Sanctum session flow, optional MFA, and then operate under a selected organization and role-scoped permissions.',
    )
    add_step_table(
        doc,
        [
            (
                'Sign in (with optional MFA)',
                'Web: /login',
                'Enter email + password. If MFA is required, enter OTP and submit again.',
            ),
            (
                'Switch active organization',
                'Top-left org selector in app layout',
                'Choose organization from dropdown to refresh project/data scope.',
            ),
            (
                'See only allowed tabs',
                'Left navigation menu',
                'Menu sections and tabs are automatically shown/hidden by permissions and project overrides.',
            ),
            (
                'Read notifications',
                'Top-right bell menu',
                'Open bell icon, review unread entries, mark single or all as read.',
            ),
        ],
    )

    add_section(
        doc,
        '2. Projects Workspace and Project Dashboard',
        'Projects page is the launch point for daily work. Users can search projects, inspect metrics, and open project dashboards.',
    )
    add_step_table(
        doc,
        [
            ('Search projects', 'Web: Projects tab (/projects)', 'Use search box, then press Search.'),
            ('Open project dashboard', 'Project card button: Open Dashboard', 'Select the target project to enter its dashboard view.'),
            (
                'Navigate to drawing viewer',
                'Web: /projects/{projectId}/drawings/{drawingId}',
                'From project dashboard, open any drawing to start snagging workflows.',
            ),
        ],
    )

    add_section(
        doc,
        '3. Drawing Viewer, Pin Plotting, and Snag Management',
        'Drawing viewer supports revision control, comparison, intelligent location assistance, barcode resolution, and bulk actions.',
    )
    add_step_table(
        doc,
        [
            (
                'Select drawing context',
                'Drawing Viewer: building/floor/revision controls',
                'Choose building, floor, current revision, and optional compare revision.',
            ),
            (
                'Create snag from pin',
                'Click on drawing canvas -> Create Snag dialog',
                'Click drawing to place normalized pin (0..1), fill fields, then create.',
            ),
            (
                'Use location suggestions',
                'Drawing Viewer intelligence panel',
                'After placing a pin, review suggested locations and confidence score.',
            ),
            (
                'Compare revisions + highlight changes',
                'Compare mode + Highlight changes switches',
                'Enable compare mode, pick second revision, then toggle change overlay.',
            ),
            (
                'Migrate pins between revisions',
                'Intelligence panel -> Pin Migration',
                'Run Preview to inspect migration, then Apply to persist migrated coordinates.',
            ),
            (
                'Resolve location by barcode/QR',
                'Scan QR/Barcode input + Open button',
                'Enter or scan barcode to deep-link into location-filtered snag view.',
            ),
            (
                'Bulk update and export snags',
                'Bulk panel above snags table',
                'Select rows, set assignee/due date, then run bulk assign/update/export.',
            ),
        ],
    )

    add_section(
        doc,
        '4. Snag Drawer - Full Collaboration Surface',
        'Snag drawer centralizes status transitions, assignment/dispatch, closeout, watchers, threaded comments with mentions, attachments, and status history.',
    )
    add_step_table(
        doc,
        [
            (
                'Transition status',
                'Snag Drawer -> Change Status panel',
                'Choose target status, add note, apply transition.',
            ),
            (
                'Assign and dispatch',
                'Snag Drawer -> Company/Team/Assignee fields',
                'Set stakeholder company/team and assignee, then Save Assignment or Dispatch.',
            ),
            (
                'Add watchers',
                'Snag Drawer -> Watchers',
                'Select member and click Watch. Remove watchers from chips if needed.',
            ),
            (
                'Threaded comments + mentions',
                'Snag Drawer -> Comments',
                'Write comment, optionally set reply parent, mention users/teams, attach files, then Add Comment.',
            ),
            (
                'Upload attachments',
                'Snag Drawer -> Attachments',
                'Choose type (photo/video/markup) and upload supporting files.',
            ),
            (
                'Review history',
                'Snag Drawer -> Status History',
                'Read full transition audit trail including actor and timestamp.',
            ),
        ],
    )

    doc.add_heading('Diagram 2 - Snag Status and Closeout Guard', level=1)
    doc.add_paragraph('Closing is guarded by closeout completion unless user has override permission.')
    doc.add_picture(str(snag_diagram), width=Inches(6.8))

    add_section(
        doc,
        '5. Closeout Templates and Closeout Instances',
        'Closeout module supports template library management and per-snag checklist execution with evidence rules.',
    )
    add_step_table(
        doc,
        [
            (
                'Manage template library',
                'Web: /templates -> Closeout tab',
                'Create, edit, clone-to-project, and delete closeout templates and checklist items.',
            ),
            (
                'Initialize closeout for snag',
                'Snag Drawer -> Closeout panel',
                'Select template and initialize/update closeout instance for the snag.',
            ),
            (
                'Complete checklist and upload evidence',
                'Snag Drawer -> checklist items',
                'Mark items complete, add notes, upload evidence files for required items.',
            ),
            (
                'Review closeout',
                'Snag Drawer -> Mark Reviewed',
                'When completion reaches 100%, reviewer can mark closeout as reviewed.',
            ),
        ],
    )

    add_section(
        doc,
        '6. Kanban Board',
        'Kanban presents snag cards by status and supports quick transitions plus deep linking into full snag details.',
    )
    add_step_table(
        doc,
        [
            ('Open board', 'Web: /board', 'Select project filter to scope board lanes.'),
            (
                'Quick move cards',
                'Buttons on each card',
                'Use available transition buttons to move snag status quickly.',
            ),
            (
                'Open full detail',
                'Click card content',
                'Open selected snag in drawer to complete richer updates.',
            ),
        ],
    )

    add_section(
        doc,
        '7. Analytics Dashboard and KPI Cards',
        'Advanced analytics include SLA metrics, Pareto/root cause views, trend/cost analysis, and custom dashboard configurations.',
    )
    add_step_table(
        doc,
        [
            ('Open analytics', 'Web: /dashboard', 'Choose project, period, and filters.'),
            (
                'Review KPI and chart groups',
                'Dashboard widgets',
                'Inspect open/overdue counts, closure timing, root causes, cost and trends.',
            ),
            (
                'Create saved dashboard config',
                'Config controls on dashboard page',
                'Select cards and filters, save new preset, update existing, or delete.',
            ),
        ],
    )

    add_section(
        doc,
        '8. Export Center',
        'Export center queues PDF/CSV/XLSX jobs and provides download once processing is completed.',
    )
    add_step_table(
        doc,
        [
            ('Request export', 'Web: /exports', 'Choose project + export type and submit request.'),
            (
                'Track status',
                'Jobs grid in Export Center',
                'Monitor queued, processing, completed, failed states.',
            ),
            (
                'Download output',
                'Download button on completed row',
                'Open generated artifact from /api/exports/{id}/download.',
            ),
        ],
    )

    add_section(
        doc,
        '9. Inspection Templates, Submissions, Approvals, Signatures, and Requests',
        'Inspections module delivers schema-driven forms, role-based approval chain, decision messaging, digital signatures, and MIR/WIR/IR request tracking.',
    )
    add_step_table(
        doc,
        [
            (
                'Build templates',
                'Web: /templates -> Inspection tab',
                'Define sections/fields and approval workflow steps, then save or clone.',
            ),
            (
                'Create submission',
                'Web: /inspections/submissions',
                'Pick template and project, create draft submission.',
            ),
            (
                'Submit and approve/reject',
                'Web: /inspections/submissions/{id}',
                'Submit draft, reviewers record decision and notes per workflow step.',
            ),
            (
                'Capture signature',
                'Submission detail -> signature area',
                'Sign required approvals and upload signature image.',
            ),
            (
                'Track MIR/WIR/IR requests',
                'Web: /inspections/requests',
                'Create request, assign, schedule, progress status to closure.',
            ),
            (
                'Inspect reports and export',
                'Web: /inspections/reports',
                'Review inspection KPIs and queue report exports.',
            ),
        ],
    )

    doc.add_heading('Diagram 3 - Inspection Approval and Signature Flow', level=1)
    doc.add_picture(str(inspection_diagram), width=Inches(6.8))

    add_section(
        doc,
        '10. Access Control, Stakeholders, Delegation, and Permission Diff',
        'Advanced RBAC supports org roles, project-level role overrides, company/team assignment structures, delegation windows, and preset comparison.',
    )
    add_step_table(
        doc,
        [
            (
                'Manage project role overrides',
                'Web: /access-control',
                'Choose project, then set role override per member and save.',
            ),
            (
                'Manage companies and teams',
                'Web: /access-control stakeholder sections',
                'Create stakeholder companies and teams for assignment workflows.',
            ),
            (
                'Set temporary delegations',
                'Web: /access-control delegation section',
                'Define from-user, to-user, date window, and scope.',
            ),
            (
                'Compare permission preset diff',
                'Web: /access-control preset comparison',
                'Select preset and compare resulting permissions against project context.',
            ),
        ],
    )

    add_section(
        doc,
        '11. Workflow Automation Engine',
        'Automation module supports conditional assignment/escalation rules, reminder policies, and recurring inspection schedules.',
    )
    add_step_table(
        doc,
        [
            ('Create automation rules', 'Web: /automation', 'Set conditions and resulting actions, then save rule.'),
            (
                'Configure reminder nudges',
                'Web: /automation reminder policies',
                'Define reminder interval and trigger scope.',
            ),
            (
                'Create recurring inspections',
                'Web: /automation recurring schedules',
                'Bind schedule to template and project cadence.',
            ),
        ],
    )

    add_section(
        doc,
        '12. Equipment and Maintenance Logs',
        'Equipment module tracks asset status, maintenance records, and linkage to snags, including optional barcode-based operations.',
    )
    add_step_table(
        doc,
        [
            ('View equipment list', 'Web: /equipment', 'Filter by project/status and open equipment details.'),
            (
                'Create equipment record',
                'Web: /equipment -> Create',
                'Enter equipment metadata and location linkage.',
            ),
            (
                'Add maintenance logs',
                'Web: equipment detail/log dialog',
                'Record maintenance action, health status, and optional snag reference.',
            ),
            (
                'Barcode flow',
                'Web and Mobile equipment tools',
                'Scan barcode to locate equipment/location instantly.',
            ),
        ],
    )

    add_section(
        doc,
        '13. Notification Preferences, Digests, and Onboarding Tours',
        'Users can tune immediate alerts, digest cadence, and quiet hours. Guided tours are persisted per user/org and can be replayed by support.',
    )
    add_step_table(
        doc,
        [
            (
                'Update notification preferences',
                'Web: /preferences/notifications',
                'Toggle channels/events and save digest frequency/timezone.',
            ),
            (
                'Run guided tours',
                'Auto-run on first entry to feature pages',
                'Tours include core, board, closeout, inspections, approvals, requests, exports, automation.',
            ),
            (
                'Replay tours or reset onboarding',
                'Web: /ops support tools',
                'Admin/support can reset onboarding state or replay selected tours for users.',
            ),
        ],
    )

    add_section(
        doc,
        '14. Ops Console and Enterprise Controls',
        'Ops console includes feature flags, usage limits, security policies, health indicators, and self-serve support operations.',
    )
    add_step_table(
        doc,
        [
            ('Feature flags', 'Web: /ops -> Feature Flags', 'Enable/disable modules per organization and project override.'),
            (
                'Usage limits',
                'Web: /ops -> Usage Limits',
                'Manage storage quota, max exports/day, and max users.',
            ),
            (
                'Security controls',
                'Web: /ops -> Security',
                'Manage MFA policy, IP allowlist, antivirus mode, and PII redaction policy.',
            ),
            (
                'Health dashboard',
                'Web: /ops -> Health',
                'Track sync errors, queue depth, and storage failures.',
            ),
            (
                'Support actions',
                'Web: /ops -> Support Tools',
                'Resend invite, reset MFA, reset onboarding, replay tours.',
            ),
        ],
    )

    add_section(
        doc,
        '15. Mobile App (Expo) - Offline-first Field Workflow',
        'Mobile app includes offline snagging, local queue, background sync policies, chunked upload, conflict resolution UI, push/deep links, and floor map navigation.',
    )
    add_step_table(
        doc,
        [
            ('Open snags list', 'Mobile tab: Snags', 'Review local snags, queue size, and run manual sync.'),
            (
                'Create snag offline',
                'Mobile: Snags -> New',
                'Fill form and queue operation. It syncs automatically later.',
            ),
            (
                'Use floor map',
                'Mobile: Snags -> Floors',
                'Navigate floors/locations quickly even when drawings are not loaded.',
            ),
            (
                'Queue comments/status/attachments',
                'Mobile: Snag Detail',
                'Add comments, queue status transitions, optimize + queue attachments.',
            ),
            (
                'Annotate attachments offline',
                'Mobile: Snag Detail -> Add + Annotate Offline',
                'Draw markups and queue both media and markup payload for sync.',
            ),
            (
                'Resolve conflicts manually',
                'Mobile tab: Conflicts',
                'Choose Use Server, Retry Local, or Merge field-by-field.',
            ),
            (
                'Tune sync windows/throttling',
                'Mobile tab: Settings -> Sync Engine',
                'Set background sync, interval, runs/minute, and active sync window.',
            ),
            (
                'Manage mobile devices and push',
                'Mobile tab: Settings',
                'Register push token, review authorized devices, revoke devices.',
            ),
        ],
    )

    doc.add_heading('Diagram 4 - Mobile Sync and Conflict Handling', level=1)
    doc.add_picture(str(mobile_diagram), width=Inches(6.8))

    add_section(
        doc,
        '16. Page and Route Index (Web)',
        'Use this index when onboarding users quickly or documenting SOPs.',
    )
    add_bullets(
        doc,
        [
            '/projects - Projects catalog and search',
            '/projects/{projectId} - Project dashboard',
            '/projects/{projectId}/drawings/{drawingId} - Drawing viewer + pin workflow',
            '/board - Snags Kanban board',
            '/dashboard - Analytics and custom KPI configs',
            '/exports - Export center',
            '/templates - Template library hub (inspection + closeout)',
            '/inspections/submissions - Inspection submissions list',
            '/inspections/submissions/{id} - Submission details, approvals, signatures, chat log',
            '/inspections/requests - MIR/WIR/IR request module',
            '/inspections/reports - Inspection reports and export queue',
            '/equipment - Equipment and maintenance logs',
            '/access-control - RBAC, stakeholders, delegation, preset diff',
            '/automation - Rules, reminders, recurring inspections',
            '/preferences/notifications - Notification and digest settings',
            '/ops - Feature flags, limits, health, security, support operations',
        ],
    )

    add_section(
        doc,
        '17. Error Handling and Troubleshooting Pointers',
        'The platform includes request correlation and structured logs for fast incident diagnosis.',
    )
    add_bullets(
        doc,
        [
            'Every API response includes X-Request-Id; keep this value when reporting issues.',
            'Frontend error messages append support reference when available.',
            'Primary logs: backend/storage/logs/api-YYYY-MM-DD.log and backend/storage/logs/laravel.log.',
            'Use Ops Health view to identify sync, queue, and storage failure trends quickly.',
        ],
    )

    doc.add_heading('Appendix - Main API Capability Groups', level=1)
    add_bullets(
        doc,
        [
            'Auth + MFA + mobile auth devices',
            'Organizations, projects, drawings, revisions, barcode location resolver',
            'Snags CRUD, transitions, dispatch, watchers, comments, attachments, escalation rules',
            'Closeout templates and closeout instances with evidence',
            'Kanban, dashboards, dashboard configs',
            'Exports and inspection report exports',
            'Inspection templates, submissions, approvals, signatures, requests, recurring schedules',
            'Automation rules and reminder policies',
            'Equipment and maintenance logs',
            'Ops controls (feature flags, limits, security, health, support)',
            'Mobile sync endpoints and chunked upload endpoints',
            'Notifications, preferences, onboarding progress',
        ],
    )

    doc.save(OUTPUT_DOCX)


if __name__ == '__main__':
    build_doc()
    print(f'Generated: {OUTPUT_DOCX}')
